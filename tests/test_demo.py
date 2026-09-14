from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from warehouse_analysis.demo import DEMO_NOTICE, build_synthetic_inputs, run_demo
from warehouse_analysis.reporting.excel_report import BUSINESS_REPORT_SHEETS


class SyntheticDemoTest(unittest.TestCase):
    def test_synthetic_inputs_are_complete_and_labeled(self) -> None:
        transactions, openings, rates = build_synthetic_inputs()
        self.assertFalse(transactions.empty)
        self.assertFalse(openings.empty)
        self.assertFalse(rates.empty)
        self.assertTrue(transactions["仓库名称"].str.startswith("演示仓").all())
        self.assertTrue(transactions["批次号"].str.startswith("DEMO-").all())
        self.assertTrue(openings["期初批次号"].str.startswith("DEMO-").all())

    def test_demo_generates_standard_input_and_business_report(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            demo = run_demo(temp_dir)
            self.assertTrue(demo.input_path.exists())
            self.assertTrue(demo.report_path.exists())
            self.assertGreater(demo.report_path.stat().st_size, 0)

            with pd.ExcelFile(demo.input_path) as workbook:
                self.assertEqual(
                    workbook.sheet_names,
                    ["演示说明", "出入库明细", "期初库存", "计价规则"],
                )
                notice = pd.read_excel(demo.input_path, sheet_name="演示说明")
                self.assertIn(DEMO_NOTICE, notice["说明"].tolist())

            styled_input = load_workbook(demo.input_path, read_only=False)
            for sheet in styled_input.worksheets:
                self.assertEqual(sheet.freeze_panes, "A2")
                self.assertFalse(sheet.sheet_view.showGridLines)
            styled_input.close()

            report = load_workbook(demo.report_path, read_only=True)
            self.assertEqual(report.sheetnames, BUSINESS_REPORT_SHEETS)
            self.assertGreater(report["批次明细"].max_row, 1)
            self.assertGreater(report["费用敏感性分析"].max_row, 1)
            details = list(report["批次明细"].iter_rows(values_only=True))
            headers = list(details[0])
            outbound_batch_index = headers.index("出库批次号")
            quantity_index = headers.index("数量（吨）")
            split_rows = [
                row for row in details[1:]
                if row[outbound_batch_index] == "DEMO-OUT-001"
            ]
            self.assertEqual(len(split_rows), 2)
            self.assertEqual(
                sorted(row[quantity_index] for row in split_rows), [30, 100]
            )
            report.close()


if __name__ == "__main__":
    unittest.main()
