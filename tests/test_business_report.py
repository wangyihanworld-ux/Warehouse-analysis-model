from __future__ import annotations

import tempfile
import unittest
from copy import deepcopy
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from pandas.testing import assert_frame_equal

from warehouse_analysis.analysis.sensitivity import (
    build_fee_sensitivity,
    build_sensitivity_ranking,
)
from warehouse_analysis.analysis.summary import build_warehouse_summary
from warehouse_analysis.core.calculation import run_engine
from warehouse_analysis.core.models import StatPeriod
from warehouse_analysis.reporting.excel_report import (
    BUSINESS_REPORT_SHEETS,
    build_business_report,
)


def _inputs() -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    transactions = pd.DataFrame(
        [
            ["测试仓库", "物料A", "合成树脂", "入库", "2025-01-02", 20, "IN-1"],
            ["测试仓库", "物料A", "合成树脂", "出库", "2025-01-03", 5, "OUT-1"],
        ],
        columns=["仓库名称", "物料名称", "品类", "方向", "日期", "数量", "批次号"],
    )
    openings = pd.DataFrame(
        [["测试仓库", "物料A", "合成树脂", 10, "2025-01-01"]],
        columns=["仓库名称", "物料名称", "品类", "期初数量", "期初批次入库基准日期"],
    )
    rates = pd.DataFrame(
        [["测试仓库", "全部", 0.7, "统一费率", "全部", "2025-01-01"]],
        columns=["仓库名称", "品类或物料", "单价", "计费方式说明", "适用层级", "生效日期"],
    )
    return transactions, openings, rates


class BusinessReportTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.temp_dir = tempfile.TemporaryDirectory()
        cls.output_path = Path(cls.temp_dir.name) / "仓储分析报告.xlsx"
        cls.inputs = _inputs()
        cls.period = StatPeriod(
            "测试期间",
            pd.Timestamp("2025-01-01"),
            pd.Timestamp("2025-01-31"),
        )
        cls.result = run_engine(*cls.inputs, [cls.period])
        cls.summary = build_warehouse_summary(
            cls.result, *cls.inputs, [cls.period]
        )
        cls.sensitivity = build_fee_sensitivity(cls.summary)
        cls.ranking = build_sensitivity_ranking(cls.sensitivity)
        cls.rate_matches = pd.DataFrame(
            [
                [
                    "测试仓库",
                    "物料A",
                    "全部",
                    "全部",
                    0.7,
                    pd.Timestamp("2025-01-01"),
                    "否",
                ]
            ],
            columns=[
                "仓库",
                "物料/品类",
                "匹配层级",
                "匹配对象",
                "单价",
                "生效日期",
                "是否人工核对",
            ],
        )
        cls.before = {
            field: deepcopy(getattr(cls.result, field))
            for field in cls.result.__dataclass_fields__
        }
        build_business_report(
            cls.result,
            cls.summary,
            cls.sensitivity,
            cls.ranking,
            pd.DataFrame(),
            cls.rate_matches,
            cls.output_path,
        )

    @classmethod
    def tearDownClass(cls) -> None:
        cls.temp_dir.cleanup()

    def test_01_excel_is_generated(self) -> None:
        self.assertTrue(self.output_path.exists())
        self.assertGreater(self.output_path.stat().st_size, 0)

    def test_02_sheet_names_are_correct(self) -> None:
        wb = load_workbook(self.output_path, read_only=True)
        self.assertEqual(set(wb.sheetnames), set(BUSINESS_REPORT_SHEETS))
        wb.close()

    def test_03_sheet_order_is_fixed(self) -> None:
        wb = load_workbook(self.output_path, read_only=True)
        self.assertEqual(wb.sheetnames, BUSINESS_REPORT_SHEETS)
        wb.close()

    def test_04_key_fields_exist(self) -> None:
        wb = load_workbook(self.output_path, read_only=True)
        home_headers = [
            cell.value for cell in next(wb["首页摘要"].iter_rows(min_row=4, max_row=4))
        ]
        self.assertIn("仓储费", home_headers)
        self.assertIn("数量守恒状态", home_headers)
        summary_headers = [
            cell.value
            for cell in next(wb["仓库经营分析"].iter_rows(min_row=1, max_row=1))
        ]
        self.assertIn("仓储费总额（元）", summary_headers)
        wb.close()

    def test_05_money_format_is_correct(self) -> None:
        wb = load_workbook(self.output_path)
        ws = wb["仓库经营分析"]
        column = [cell.value for cell in ws[1]].index("仓储费总额（元）") + 1
        self.assertEqual(ws.cell(2, column).number_format, "#,##0.00")
        wb.close()

    def test_06_quantity_format_is_correct(self) -> None:
        wb = load_workbook(self.output_path)
        ws = wb["仓库经营分析"]
        column = [cell.value for cell in ws[1]].index("期末库存量（吨）") + 1
        self.assertEqual(ws.cell(2, column).number_format, "#,##0.000")
        wb.close()

    def test_07_all_sheets_have_freeze_panes_and_filter(self) -> None:
        wb = load_workbook(self.output_path)
        for ws in wb.worksheets:
            self.assertIsNotNone(ws.freeze_panes, ws.title)
            self.assertIsNotNone(ws.auto_filter.ref, ws.title)
        wb.close()

    def test_08_excel_can_be_reopened(self) -> None:
        wb = load_workbook(self.output_path, data_only=False)
        self.assertEqual(len(wb.worksheets), 12)
        wb.close()

    def test_09_report_data_matches_dataframes(self) -> None:
        wb = load_workbook(self.output_path, data_only=True)
        ws = wb["仓库经营分析"]
        headers = [cell.value for cell in ws[1]]
        values = [cell.value for cell in ws[2]]
        report_row = dict(zip(headers, values))
        expected = self.summary.iloc[0]
        self.assertEqual(report_row["仓库名称"], expected["仓库名称"])
        self.assertAlmostEqual(
            report_row["仓储费总额（元）"], expected["仓储费总额（元）"]
        )
        ws = wb["费用敏感性分析"]
        self.assertEqual(ws.max_row - 1, len(self.sensitivity))
        wb.close()

    def test_10_report_generation_does_not_mutate_core_result(self) -> None:
        for field, before in self.before.items():
            assert_frame_equal(getattr(self.result, field), before)


if __name__ == "__main__":
    unittest.main()
