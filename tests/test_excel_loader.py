from __future__ import annotations

import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

import pandas as pd
from openpyxl import Workbook

from warehouse_analysis.core.rate_engine import RateResolver
from warehouse_analysis.core.validation import normalize_inputs
from warehouse_analysis.io.excel_loader import (
    InputWorkbookError,
    load_analysis_inputs,
    load_business_workbook,
)


TRANSACTION_HEADERS = [
    "仓库名称", "物料名称", "品类", "方向", "日期", "数量", "批次号"
]
OPENING_HEADERS = [
    "仓库名称", "物料名称", "品类", "期初数量", "期初批次入库基准日期"
]
RATE_HEADERS = [
    "仓库名称", "品类或物料", "单价", "计费方式说明", "适用层级", "生效日期"
]


def add_sheet(wb: Workbook, name: str, headers: list[str], row: list) -> None:
    ws = wb.create_sheet(name)
    ws.append(headers)
    ws.append(row)


def build_business_file(
    path: Path,
    *,
    opening_sheet: str = "期初库存",
    include_rate: bool = False,
    include_both_opening_sheets: bool = False,
    transaction_headers: list[str] | None = None,
) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    add_sheet(
        wb,
        "出入库明细",
        transaction_headers or TRANSACTION_HEADERS,
        ["测试仓库", "物料A", "合成树脂", "入库", "2025-01-02", 10, "B1"][
            : len(transaction_headers or TRANSACTION_HEADERS)
        ],
    )
    opening_row = ["测试仓库", "物料A", "合成树脂", 5, "2025-01-01"]
    add_sheet(wb, opening_sheet, OPENING_HEADERS, opening_row)
    if include_both_opening_sheets:
        other = "期初存货" if opening_sheet == "期初库存" else "期初库存"
        add_sheet(wb, other, OPENING_HEADERS, opening_row)
    if include_rate:
        add_sheet(
            wb,
            "计价规则",
            RATE_HEADERS,
            ["测试仓库", "全部", 0.7, "统一费率", "全部", "2025-01-01"],
        )
    wb.save(path)


def build_rate_file(path: Path, sheet_name: str = "计价规则") -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(RATE_HEADERS)
    ws.append(["测试仓库", "全部", 0.7, "统一费率", "全部", "2025-01-01"])
    wb.save(path)


class ExcelLoaderStageThreeTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.base = Path(self.temp.name)

    def tearDown(self) -> None:
        self.temp.cleanup()

    def test_independent_rate_config_is_loaded(self) -> None:
        business = self.base / "business.xlsx"
        rates = self.base / "rates.xlsx"
        build_business_file(business)
        build_rate_file(rates)

        tx, op, rt = load_analysis_inputs(business, rates)

        self.assertEqual(len(tx), 1)
        self.assertEqual(len(op), 1)
        self.assertEqual(len(rt), 1)
        self.assertEqual(rt.iloc[0]["单价"], 0.7)

    def test_combined_three_sheet_workbook_is_compatible(self) -> None:
        combined = self.base / "combined.xlsx"
        build_business_file(
            combined, opening_sheet="期初存货", include_rate=True
        )

        tx, op, rt = load_analysis_inputs(combined)

        self.assertEqual((len(tx), len(op), len(rt)), (1, 1, 1))

    def test_opening_inventory_sheet_is_preferred_name(self) -> None:
        business = self.base / "business.xlsx"
        build_business_file(business, opening_sheet="期初库存")

        _, openings = load_business_workbook(business)

        self.assertEqual(openings.iloc[0]["期初数量"], 5)

    def test_opening_stock_legacy_sheet_is_supported(self) -> None:
        business = self.base / "business.xlsx"
        build_business_file(business, opening_sheet="期初存货")

        _, openings = load_business_workbook(business)

        self.assertEqual(openings.iloc[0]["期初数量"], 5)

    def test_two_opening_sheets_raise_clear_error(self) -> None:
        business = self.base / "business.xlsx"
        build_business_file(
            business,
            opening_sheet="期初库存",
            include_both_opening_sheets=True,
        )

        with self.assertRaises(InputWorkbookError) as captured:
            load_business_workbook(business)

        message = str(captured.exception)
        self.assertIn("同时存在", message)
        self.assertIn("请只保留一个期初工作表", message)

    def test_missing_field_error_contains_file_sheet_and_suggestion(self) -> None:
        business = self.base / "missing.xlsx"
        bad_headers = [h for h in TRANSACTION_HEADERS if h != "数量"]
        build_business_file(business, transaction_headers=bad_headers)

        with self.assertRaises(InputWorkbookError) as captured:
            load_business_workbook(business)

        message = str(captured.exception)
        self.assertIn(f"文件：{business}", message)
        self.assertIn("Sheet：出入库明细", message)
        self.assertIn("缺少字段：数量", message)
        self.assertIn("请补充数量列", message)
        self.assertIsInstance(captured.exception.__cause__, ValueError)

    def test_independent_and_combined_rates_match_identically(self) -> None:
        combined = self.base / "combined.xlsx"
        business = self.base / "business.xlsx"
        rates = self.base / "rates.xlsx"
        build_business_file(
            combined, opening_sheet="期初存货", include_rate=True
        )
        build_business_file(business)
        build_rate_file(rates)

        combined_frames = normalize_inputs(*load_analysis_inputs(combined))
        separate_frames = normalize_inputs(
            *load_analysis_inputs(business, rates)
        )
        old_resolver = RateResolver(
            combined_frames[2], {"物料A"}, {"合成树脂"}
        )
        new_resolver = RateResolver(
            separate_frames[2], {"物料A"}, {"合成树脂"}
        )
        old_match = old_resolver.resolve(
            "测试仓库", "物料A", "合成树脂", pd.Timestamp("2025-01-02")
        )
        new_match = new_resolver.resolve(
            "测试仓库", "物料A", "合成树脂", pd.Timestamp("2025-01-02")
        )

        self.assertEqual(tuple(new_match), tuple(old_match))
        self.assertEqual(new_match.rate, Decimal("0.7"))
        self.assertEqual(new_match.level, old_match.level)


if __name__ == "__main__":
    unittest.main()
