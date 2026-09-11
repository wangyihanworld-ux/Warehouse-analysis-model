from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from warehouse_analysis.io.excel_loader import (
    OPENING_COLUMNS,
    RATE_COLUMNS,
    TRANSACTION_COLUMNS,
)


HEADER_FILL = "4472C4"
HEADER_FONT = "FFFFFF"
NOTE_FILL = "D9EAF7"
INPUT_FILL = "FFF2CC"
GRID = Side(style="thin", color="B7C9D6")
FONT_NAME = "微软雅黑"

INPUT_TEMPLATE_NAME = "仓储分析输入模板.xlsx"
RATE_TEMPLATE_NAME = "rate_config_template.xlsx"

EXTENDED_RATE_COLUMNS = RATE_COLUMNS + ["适用层级", "生效日期", "是否自动计算", "特殊规则类型"]


def _style_header(ws, headers: list[str]) -> None:
    """设置模板表头；仅影响展示，不影响任何计算模块。"""
    for column, header in enumerate(headers, 1):
        cell = ws.cell(1, column, header)
        cell.font = Font(name=FONT_NAME, bold=True, color=HEADER_FONT)
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=GRID)
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    ws.sheet_view.showGridLines = False


def _set_widths(ws, widths: dict[str, float]) -> None:
    for column, width in widths.items():
        ws.column_dimensions[column].width = width


def _add_instruction_sheet(
    wb: Workbook,
    rows: list[tuple[str, str, str, str, str]],
) -> None:
    ws = wb.create_sheet("字段说明")
    headers = ["工作表", "字段", "是否必填", "填写说明", "示例"]
    _style_header(ws, headers)
    for row in rows:
        ws.append(row)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name=FONT_NAME, size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=GRID)
        row[0].fill = PatternFill("solid", fgColor=NOTE_FILL)
    _set_widths(
        ws,
        {"A": 16, "B": 24, "C": 12, "D": 52, "E": 24},
    )


def create_input_template(path: Path) -> None:
    """生成业务数据输入模板，不读取或修改计算模块。"""
    wb = Workbook()
    tx = wb.active
    tx.title = "出入库明细"
    _style_header(tx, TRANSACTION_COLUMNS)
    _set_widths(
        tx,
        {"A": 20, "B": 28, "C": 18, "D": 12, "E": 15, "F": 14, "G": 22},
    )
    tx["E2"].number_format = "yyyy-mm-dd"
    tx["F2"].number_format = "#,##0.000"
    direction_validation = DataValidation(
        type="list", formula1='"入库,出库"', allow_blank=True
    )
    direction_validation.error = "方向只能填写“入库”或“出库”。"
    direction_validation.errorTitle = "方向填写错误"
    direction_validation.prompt = "请选择入库或出库。"
    direction_validation.promptTitle = "方向"
    direction_validation.showErrorMessage = True
    direction_validation.showInputMessage = True
    tx.add_data_validation(direction_validation)
    direction_validation.add("D2:D10000")

    opening = wb.create_sheet("期初库存")
    _style_header(opening, OPENING_COLUMNS)
    _set_widths(
        opening,
        {"A": 20, "B": 28, "C": 18, "D": 16, "E": 26},
    )
    opening["D2"].number_format = "#,##0.000"
    opening["E2"].number_format = "yyyy-mm-dd"

    instructions = [
        ("出入库明细", "仓库名称", "是", "库存所属仓库的标准名称。", "四川力庆库"),
        ("出入库明细", "物料名称", "是", "物料的标准名称，同一物料应保持完全一致。", "独山子HD5420GA"),
        ("出入库明细", "品类", "是", "用于费率匹配和品类汇总。", "合成树脂"),
        ("出入库明细", "方向", "是", "只能填写“入库”或“出库”。", "入库"),
        ("出入库明细", "日期", "是", "业务发生日期，格式YYYY-MM-DD。", "2025-01-02"),
        ("出入库明细", "数量", "是", "正数，单位为吨；方向由“方向”字段决定。", "32.500"),
        ("出入库明细", "批次号", "是", "业务批次或单据明细的可追溯编号。", "B20250102001"),
        ("期初库存", "仓库名称", "是", "库存所属仓库的标准名称。", "四川力庆库"),
        ("期初库存", "物料名称", "是", "必须与流水中的物料名称保持一致。", "独山子HD5420GA"),
        ("期初库存", "品类", "是", "必须与物料对应品类保持一致。", "合成树脂"),
        ("期初库存", "期初数量", "是", "统计期开始前已有库存，正数，单位为吨。", "50.000"),
        (
            "期初库存",
            "期初批次入库基准日期",
            "是",
            "作为期初库存计算库龄的入库基准日期，格式YYYY-MM-DD。",
            "2025-01-01",
        ),
    ]
    _add_instruction_sheet(wb, instructions)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def create_rate_template(path: Path) -> None:
    """生成独立费率配置模板，不改变现有费率匹配算法。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "计价规则"
    _style_header(ws, EXTENDED_RATE_COLUMNS)
    _set_widths(
        ws,
        {"A": 20, "B": 30, "C": 14, "D": 34, "E": 14, "F": 15, "G": 18, "H": 22},
    )
    ws.append(
        ["示例仓库", "全部", 0.7, "统一单价", "全部", date(2025, 1, 1), "是", None]
    )
    ws.append(
        [
            "示例仓库",
            "示例物料",
            0,
            "货转货物延续上家费率",
            "物料",
            date(2025, 1, 1),
            "否",
            "人工核对",
        ]
    )
    for row in ws.iter_rows(min_row=2, max_row=3):
        for cell in row:
            cell.font = Font(name=FONT_NAME, size=10)
            cell.border = Border(bottom=GRID)
        row[2].number_format = "0.0000"
        row[5].number_format = "yyyy-mm-dd"
        row[6].fill = PatternFill("solid", fgColor=INPUT_FILL)
        row[7].fill = PatternFill("solid", fgColor=INPUT_FILL)

    level_validation = DataValidation(
        type="list", formula1='"物料,品类,全部"', allow_blank=False
    )
    auto_validation = DataValidation(
        type="list", formula1='"是,否"', allow_blank=False
    )
    ws.add_data_validation(level_validation)
    ws.add_data_validation(auto_validation)
    level_validation.add("E2:E10000")
    auto_validation.add("G2:G10000")

    instructions = [
        ("计价规则", "仓库名称", "是", "规则适用的仓库标准名称。", "四川力庆库"),
        (
            "计价规则",
            "品类或物料",
            "是",
            "适用层级为物料时填物料名称；品类时填品类；全部时填“全部”。",
            "全部",
        ),
        (
            "计价规则",
            "单价",
            "是",
            "单位为元/吨/天；当前模板中的人工规则示例填0，后续由人工核对标记控制。",
            "0.7",
        ),
        ("计价规则", "计费方式说明", "是", "描述合同计费规则。", "统一单价"),
        ("计价规则", "适用层级", "是", "只能填写物料、品类或全部。", "全部"),
        ("计价规则", "生效日期", "是", "费率生效日期，格式YYYY-MM-DD。", "2025-01-01"),
        (
            "计价规则",
            "是否自动计算",
            "是",
            "普通规则填“是”；无法自动判断的规则填“否”。",
            "是",
        ),
        (
            "计价规则",
            "特殊规则类型",
            "否",
            "是否自动计算为“否”时填写，如“人工核对”。",
            "人工核对",
        ),
    ]
    _add_instruction_sheet(wb, instructions)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def generate_templates(output_dir: Path) -> tuple[Path, Path]:
    """生成并重新打开验证两个模板，返回文件路径。"""
    output_dir = Path(output_dir)
    input_path = output_dir / INPUT_TEMPLATE_NAME
    rate_path = output_dir / RATE_TEMPLATE_NAME
    create_input_template(input_path)
    create_rate_template(rate_path)
    for path in (input_path, rate_path):
        workbook = load_workbook(path, read_only=False, data_only=False)
        workbook.close()
    return input_path, rate_path


if __name__ == "__main__":
    generate_templates(Path.cwd())
