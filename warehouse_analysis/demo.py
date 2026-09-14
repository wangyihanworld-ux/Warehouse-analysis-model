"""完全脱敏的仓储分析演示入口。

运行 ``python -m warehouse_analysis.demo`` 会创建合成 Excel 输入文件，并使用
与正式命令行相同的服务层生成一份 FIFO、仓储费、库龄和敏感性分析报告。
本模块中的仓库、物料、批次、数量与费率均为人为构造，不读取任何外部业务文件。
"""

from __future__ import annotations

import argparse
import json
from dataclasses import dataclass
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from warehouse_analysis.core.models import StatPeriod
from warehouse_analysis.service import AnalysisRunResult, run_analysis


DEMO_NOTICE = "本工作簿仅含确定性合成演示数据，不含真实业务、客户或合同信息。"
DEFAULT_OUTPUT_DIR = Path("demo_output")
DEMO_PERIOD = StatPeriod(
    "2025年第一季度（合成演示）",
    pd.Timestamp("2025-01-01"),
    pd.Timestamp("2025-03-31"),
)


@dataclass(frozen=True)
class DemoRun:
    """一次演示运行生成的输入与报告路径。"""

    input_path: Path
    report_path: Path
    summary: dict[str, object]


def build_synthetic_inputs() -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """返回固定的合成流水、期初库存与费率，不访问任何外部数据源。"""
    transactions = pd.DataFrame(
        [
            ["演示仓A", "示例树脂A", "合成材料", "入库", "2025-01-08", 120, "DEMO-IN-001"],
            ["演示仓A", "示例树脂B", "合成材料", "入库", "2025-01-20", 80, "DEMO-IN-002"],
            ["演示仓A", "示例树脂A", "合成材料", "出库", "2025-02-05", 130, "DEMO-OUT-001"],
            ["演示仓A", "示例树脂B", "合成材料", "入库", "2025-02-18", 60, "DEMO-IN-003"],
            ["演示仓A", "示例树脂B", "合成材料", "出库", "2025-03-12", 70, "DEMO-OUT-002"],
            ["演示仓B", "示例包装膜", "包装材料", "入库", "2025-01-15", 150, "DEMO-IN-004"],
            ["演示仓B", "示例包装膜", "包装材料", "出库", "2025-02-28", 65, "DEMO-OUT-003"],
            ["演示仓B", "示例包装膜", "包装材料", "入库", "2025-03-10", 40, "DEMO-IN-005"],
        ],
        columns=["仓库名称", "物料名称", "品类", "方向", "日期", "数量", "批次号"],
    )
    openings = pd.DataFrame(
        [
            ["演示仓A", "示例树脂A", "合成材料", 100, "2024-12-10", "DEMO-OPEN-001"],
            ["演示仓B", "示例包装膜", "包装材料", 90, "2024-12-18", "DEMO-OPEN-002"],
        ],
        columns=[
            "仓库名称", "物料名称", "品类", "期初数量",
            "期初批次入库基准日期", "期初批次号",
        ],
    )
    rates = pd.DataFrame(
        [
            ["演示仓A", "合成材料", 0.58, "合成演示品类日费率", "品类", "2025-01-01"],
            ["演示仓B", "包装材料", 0.42, "合成演示品类日费率", "品类", "2025-01-01"],
        ],
        columns=["仓库名称", "品类或物料", "单价", "计费方式说明", "适用层级", "生效日期"],
    )
    return transactions, openings, rates


def _unique_path(directory: Path, stem: str) -> Path:
    """生成不覆盖既有演示文件的 xlsx 路径。"""
    candidate = directory / f"{stem}.xlsx"
    for index in range(1, 1000):
        if not candidate.exists():
            return candidate
        candidate = directory / f"{stem}_{index:03d}.xlsx"
    raise FileExistsError("演示文件数量超过 999，无法创建新的演示文件。")


def write_synthetic_workbook(output_dir: str | Path) -> Path:
    """将合成三表和数据声明写入一个标准输入工作簿。"""
    directory = Path(output_dir)
    directory.mkdir(parents=True, exist_ok=True)
    input_path = _unique_path(directory, "合成演示输入数据")
    transactions, openings, rates = build_synthetic_inputs()
    notice = pd.DataFrame(
        {
            "说明": [
                DEMO_NOTICE,
                "仓库、物料、批次、数量和费率均为示例，不对应任何客户或合同。",
                "可用本文件验证 FIFO、仓储费用、库龄、月末库存和敏感性分析报告。",
            ]
        }
    )
    with pd.ExcelWriter(input_path, engine="openpyxl") as writer:
        notice.to_excel(writer, sheet_name="演示说明", index=False)
        transactions.to_excel(writer, sheet_name="出入库明细", index=False)
        openings.to_excel(writer, sheet_name="期初库存", index=False)
        rates.to_excel(writer, sheet_name="计价规则", index=False)
    _style_synthetic_workbook(input_path)
    return input_path


def _style_synthetic_workbook(path: Path) -> None:
    """设置演示输入格式，便于逐表理解和检查。"""
    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="4472C4")
    for sheet in workbook.worksheets:
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 28
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = Font(name="微软雅黑", bold=True, color="FFFFFF")
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="微软雅黑", size=10)
                cell.alignment = Alignment(vertical="center")
        for column in sheet.columns:
            letter = column[0].column_letter
            width = max(len(str(cell.value or "")) for cell in column) + 3
            sheet.column_dimensions[letter].width = min(max(width, 12), 42)
    workbook["演示说明"].column_dimensions["A"].width = 78
    workbook.save(path)


def run_demo(output_dir: str | Path = DEFAULT_OUTPUT_DIR) -> DemoRun:
    """生成合成输入并调用正式服务层输出完整示例报告。"""
    directory = Path(output_dir)
    input_path = write_synthetic_workbook(directory)
    result: AnalysisRunResult = run_analysis(
        input_path,
        rate_config=None,
        period=DEMO_PERIOD,
        output_path=directory,
    )
    if not result.success or result.output_path is None:
        raise RuntimeError(result.error or "合成演示报告生成失败")
    return DemoRun(
        input_path=input_path,
        report_path=result.output_path,
        summary=result.summary,
    )


def main() -> None:
    parser = argparse.ArgumentParser(description="生成完全脱敏的 FIFO 仓储分析演示")
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=DEFAULT_OUTPUT_DIR,
        help="合成输入和示例报告的输出目录（默认：demo_output）",
    )
    args = parser.parse_args()
    demo = run_demo(args.output_dir)
    print(
        json.dumps(
            {
                "运行状态": "成功",
                "数据声明": DEMO_NOTICE,
                "合成输入文件": str(demo.input_path),
                "示例报告": str(demo.report_path),
                **demo.summary,
            },
            ensure_ascii=False,
            indent=2,
            default=str,
        )
    )


if __name__ == "__main__":
    main()
