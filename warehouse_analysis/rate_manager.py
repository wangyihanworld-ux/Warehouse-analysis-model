from __future__ import annotations

from pathlib import Path
import sys
from typing import Any, Mapping
from uuid import uuid4

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


PROJECT_ROOT = (
    Path(sys.executable).resolve().parent
    if getattr(sys, "frozen", False)
    else Path(__file__).resolve().parents[1]
)
DEFAULT_RATE_DATABASE = PROJECT_ROOT / "config" / "rate_database.xlsx"
RATE_DATABASE_SHEET = "费率数据库"
RATE_DATABASE_COLUMNS = [
    "仓库编号",
    "仓库名称",
    "仓库别名",
    "计费规则",
    "计价方式",
    "品类",
    "天数下限",
    "天数上限",
    "单价",
    "生效日期",
    "失效日期",
    "是否启用",
    "备注",
    "选择项",
]
LEGACY_COLUMNS = [
    "仓库编号",
    "仓库名称",
    "计价方式",
    "品类",
    "天数下限",
    "天数上限",
    "单价",
]
EDITABLE_FIELDS = {"单价", "生效日期", "计费规则", "备注"}


class RateRuleCompatibilityError(ValueError):
    """费率规则存在，但冻结费用核心不能自动执行该规则。"""

    def __init__(self, diagnosis: Mapping[str, object]):
        self.diagnosis = dict(diagnosis)
        super().__init__(
            f"规则“{self.diagnosis.get('规则名称', '')}”需要业务确认："
            f"{self.diagnosis.get('为什么无法自动计算', '')}"
        )


def _selection(warehouse: object, rule: object) -> str:
    return f"{str(warehouse).strip()}+{str(rule).strip()}"


def _normalize_database(frame: pd.DataFrame) -> pd.DataFrame:
    if "仓库别名" not in frame.columns:
        frame = frame.copy()
        frame["仓库别名"] = ""
    missing = [
        column for column in RATE_DATABASE_COLUMNS if column not in frame.columns
    ]
    if missing:
        raise ValueError(f"费率数据库缺少字段：{', '.join(missing)}")
    result = frame[RATE_DATABASE_COLUMNS].copy()
    for column in ("仓库编号", "仓库名称", "计费规则", "计价方式", "品类"):
        if result[column].isna().any():
            raise ValueError(f"费率数据库字段“{column}”存在空值")
        result[column] = result[column].astype(str).str.strip()
    result["仓库别名"] = result["仓库别名"].fillna("").astype(str).str.strip()
    result["单价"] = pd.to_numeric(result["单价"], errors="raise")
    if (result["单价"] < 0).any():
        raise ValueError("费率数据库单价不能小于0")
    result["生效日期"] = pd.to_datetime(
        result["生效日期"], errors="raise"
    ).dt.normalize()
    result["失效日期"] = pd.to_datetime(
        result["失效日期"], errors="coerce"
    ).dt.normalize()
    result["是否启用"] = (
        result["是否启用"].fillna("是").astype(str).str.strip()
    )
    invalid_enabled = sorted(set(result["是否启用"]) - {"是", "否"})
    if invalid_enabled:
        raise ValueError(f"是否启用只能填写是或否：{invalid_enabled}")
    result["备注"] = result["备注"].fillna("").astype(str)
    result["选择项"] = [
        _selection(warehouse, rule)
        for warehouse, rule in zip(result["仓库名称"], result["计费规则"])
    ]
    return result


def _legacy_to_database(
    frame: pd.DataFrame,
    *,
    effective_date: str | pd.Timestamp = "2025-01-01",
    note: str = "由rate_config.xlsx导入",
) -> pd.DataFrame:
    missing = [column for column in LEGACY_COLUMNS if column not in frame.columns]
    if missing:
        raise ValueError(f"导入费率配置缺少字段：{', '.join(missing)}")
    result = frame[LEGACY_COLUMNS].copy()
    result.insert(2, "计费规则", result["计价方式"].astype(str).str.strip())
    result["仓库别名"] = ""
    result["生效日期"] = pd.Timestamp(effective_date)
    result["失效日期"] = pd.NaT
    result["是否启用"] = "是"
    result["备注"] = note
    result["选择项"] = [
        _selection(warehouse, rule)
        for warehouse, rule in zip(result["仓库名称"], result["计费规则"])
    ]
    return _normalize_database(result)


def _read_rate_file(path: Path) -> pd.DataFrame:
    with pd.ExcelFile(path) as workbook:
        if RATE_DATABASE_SHEET in workbook.sheet_names:
            sheet = RATE_DATABASE_SHEET
        elif "费率配置" in workbook.sheet_names:
            sheet = "费率配置"
        elif len(workbook.sheet_names) == 1:
            sheet = workbook.sheet_names[0]
        else:
            raise ValueError("费率文件中未找到“费率数据库”或“费率配置”Sheet")
    return pd.read_excel(path, sheet_name=sheet)


def _write_database(frame: pd.DataFrame, path: Path) -> None:
    """原子写入费率库，避免写入中断损坏正式文件。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    normalized = _normalize_database(frame)
    temporary = path.with_name(f".{path.stem}.{uuid4().hex}.tmp.xlsx")
    try:
        with pd.ExcelWriter(temporary, engine="openpyxl") as writer:
            normalized.to_excel(
                writer, sheet_name=RATE_DATABASE_SHEET, index=False
            )
        workbook = load_workbook(temporary)
        sheet = workbook[RATE_DATABASE_SHEET]
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.font = Font(name="微软雅黑", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="4472C4")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for column_index, header in enumerate(RATE_DATABASE_COLUMNS, 1):
            values = [
                str(sheet.cell(row, column_index).value or "")
                for row in range(1, min(sheet.max_row, 100) + 1)
            ]
            sheet.column_dimensions[get_column_letter(column_index)].width = min(
                max(max(map(len, values), default=8) + 2, 12), 34
            )
            for row in range(2, sheet.max_row + 1):
                cell = sheet.cell(row, column_index)
                if header in {"生效日期", "失效日期"}:
                    cell.number_format = "yyyy-mm-dd"
                elif header == "单价":
                    cell.number_format = "#,##0.0000"
                elif header in {"天数下限", "天数上限"}:
                    cell.number_format = "#,##0"
        workbook.save(temporary)
        temporary.replace(path)
    finally:
        if temporary.exists():
            temporary.unlink()


def initialize_rate_database(
    source_path: str | Path,
    database_path: str | Path = DEFAULT_RATE_DATABASE,
) -> Path:
    """从旧rate_config.xlsx初始化费率库；拒绝覆盖已存在数据库。"""
    source = Path(source_path)
    destination = Path(database_path)
    if destination.exists():
        raise FileExistsError(f"费率数据库已存在，不允许覆盖：{destination}")
    imported = _legacy_to_database(_read_rate_file(source))
    _write_database(imported, destination)
    return destination


def load_rate_database(
    database_path: str | Path = DEFAULT_RATE_DATABASE,
) -> pd.DataFrame:
    """读取并验证系统费率库。"""
    path = Path(database_path)
    if not path.is_file():
        raise FileNotFoundError(f"费率数据库不存在：{path}")
    return _normalize_database(_read_rate_file(path))


def search_rate_rule(
    keyword: str,
    database_path: str | Path = DEFAULT_RATE_DATABASE,
    *,
    include_disabled: bool = False,
) -> pd.DataFrame:
    """按仓库名称或计费规则模糊搜索，返回唯一选择项。"""
    data = load_rate_database(database_path)
    if not include_disabled:
        data = data[data["是否启用"] == "是"]
    term = str(keyword).strip()
    if term:
        mask = data["仓库名称"].str.contains(
            term, case=False, regex=False
        ) | data["计费规则"].str.contains(term, case=False, regex=False)
        data = data[mask]
    result = (
        data.sort_values(["生效日期"], ascending=False, kind="stable")
        .drop_duplicates("选择项")
        [["选择项", "仓库编号", "仓库名称", "计费规则", "是否启用"]]
        .sort_values(["仓库名称", "计费规则"], kind="stable")
        .reset_index(drop=True)
    )
    return result


def add_warehouse_alias(
    warehouse: str,
    alias: str,
    database_path: str | Path = DEFAULT_RATE_DATABASE,
) -> int:
    """为正式仓库名追加别名；不改正式名称、不覆盖规则历史。"""
    formal = str(warehouse).strip()
    alias_name = str(alias).strip()
    if not formal or not alias_name:
        raise ValueError("仓库正式名称和别名不能为空")
    data = load_rate_database(database_path)
    mask = data["仓库名称"] == formal
    if not mask.any():
        raise ValueError(f"费率数据库不存在仓库：{formal}")
    changed = 0
    for index in data.index[mask]:
        existing = [
            item.strip()
            for item in str(data.at[index, "仓库别名"] or "").replace("；", ";").split(";")
            if item.strip() and item.strip().lower() != "nan"
        ]
        if alias_name not in existing:
            existing.append(alias_name)
            data.at[index, "仓库别名"] = "；".join(existing)
            changed += 1
    if changed:
        _write_database(data, Path(database_path))
    return changed


def rate_match_diagnosis(
    warehouse: str,
    material: str,
    category: str,
    stat_date: str | pd.Timestamp,
    database_path: str | Path = DEFAULT_RATE_DATABASE,
) -> dict[str, object]:
    """诊断费率匹配失败原因，不参与核心费率选择或费用计算。"""
    data = load_rate_database(database_path)
    date = pd.Timestamp(stat_date).normalize()
    warehouse_name = str(warehouse).strip()
    material_name = str(material).strip()
    category_name = str(category).strip()
    warehouse_rules = data[data["仓库名称"] == warehouse_name].copy()
    base = {
        "仓库名称": warehouse_name,
        "物料名称": material_name,
        "品类": category_name,
        "统计日期": date.strftime("%Y-%m-%d"),
        "业务日期": date.strftime("%Y-%m-%d"),
        "仓库是否存在": "是" if not warehouse_rules.empty else "否",
        "仓库已有规则数量": int(len(warehouse_rules)),
        "候选规则数量": int(len(warehouse_rules)),
        "当前规则": warehouse_rules[
            ["选择项", "品类", "单价", "生效日期", "失效日期", "是否启用"]
        ].to_dict("records"),
    }
    empty_search = {"物料级": "无", "品类级": "无", "仓库级": "无"}
    if warehouse_rules.empty:
        return {
            **base,
            "当前搜索": empty_search,
            "匹配状态": "失败",
            "失败原因": "没有该仓库规则",
            "建议": "请进入费率管理新增该仓库规则。",
        }
    enabled = warehouse_rules[warehouse_rules["是否启用"] == "是"].copy()
    started = enabled[enabled["生效日期"] <= date].copy()
    active = started[
        started["失效日期"].isna() | (started["失效日期"] >= date)
    ].copy()
    search = {
        "物料级": int((active["品类"].astype(str).str.strip() == material_name).sum()),
        "品类级": int((active["品类"].astype(str).str.strip() == category_name).sum()),
        "仓库级": int((active["品类"].astype(str).str.strip() == "全部").sum()),
    }
    search = {level: (count if count else "无") for level, count in search.items()}
    if active.empty:
        if enabled.empty:
            reason = "仓库规则已停用"
            suggestion = "请确认停用原因；如需恢复使用，应由业务人员维护启用状态。"
        elif started.empty:
            next_date = enabled["生效日期"].min().strftime("%Y-%m-%d")
            reason = f"仓库规则尚未生效；最早生效日期为{next_date}"
            suggestion = "请新增覆盖统计日期的规则版本，或核对业务统计日期。"
        else:
            last_date = started["失效日期"].dropna().max()
            last_text = (
                last_date.strftime("%Y-%m-%d")
                if pd.notna(last_date)
                else "<空>"
            )
            reason = f"仓库规则日期失效；最近失效日期为{last_text}"
            suggestion = "请新增费率版本或核对失效日期；系统不会自动修改日期。"
        return {
            **base,
            "当前搜索": empty_search,
            "搜索层级": ["物料级", "品类级", "仓库级"],
            "匹配状态": "失败",
            "失败原因": reason,
            "建议": suggestion,
            "建议动作": suggestion,
        }
    priority_values = [
        (material_name, 3),
        (category_name, 2),
        ("全部", 1),
    ]
    matched = pd.DataFrame()
    matched_level = ""
    for value, priority in priority_values:
        if not value:
            continue
        candidate = active[active["品类"].astype(str).str.strip() == value]
        if not candidate.empty:
            matched = candidate
            matched_level = {3: "物料", 2: "品类", 1: "仓库全部"}[priority]
            break
    if matched.empty:
        return {
            **base,
            "当前搜索": search,
            "搜索层级": ["物料级", "品类级", "仓库级"],
            "匹配状态": "失败",
            "失败原因": "有有效仓库规则，但品类或物料不匹配",
            "建议": f"请新增品类“{category_name}”或物料“{material_name}”规则。",
            "建议动作": f"请新增品类“{category_name}”或物料“{material_name}”规则。",
        }
    latest = matched["生效日期"].max()
    latest_rows = matched[matched["生效日期"] == latest]
    comparable = latest_rows[
        ["选择项", "品类", "天数下限", "天数上限", "单价"]
    ].fillna("<空>")
    if len(comparable.drop_duplicates()) > 1:
        return {
            **base,
            "当前搜索": search,
            "匹配状态": "失败",
            "失败原因": "存在多个同层级有效规则冲突",
            "匹配层级": matched_level,
            "建议": "请停用重复规则，或明确不同规则的适用期间。",
        }
    selected = latest_rows.iloc[0]
    return {
        **base,
        "当前搜索": search,
        "搜索层级": ["物料级", "品类级", "仓库级"],
        "匹配状态": "成功",
        "失败原因": "",
        "匹配层级": matched_level,
        "匹配规则": str(selected["选择项"]),
        "匹配单价": float(selected["单价"]),
        "规则生效日期": pd.Timestamp(selected["生效日期"]).strftime(
            "%Y-%m-%d"
        ),
        "建议": "费率可以匹配。",
        "建议动作": "费率可以匹配。",
    }


def add_rate_rule(
    record: Mapping[str, Any],
    database_path: str | Path = DEFAULT_RATE_DATABASE,
) -> pd.DataFrame:
    """新增一条费率明细；相同版本和档位已存在时拒绝写入。"""
    path = Path(database_path)
    data = (
        load_rate_database(path)
        if path.exists()
        else pd.DataFrame(columns=RATE_DATABASE_COLUMNS)
    )
    row = {column: record.get(column) for column in RATE_DATABASE_COLUMNS}
    row["失效日期"] = record.get("失效日期", pd.NaT)
    row["是否启用"] = record.get("是否启用", "是")
    row["备注"] = record.get("备注", "")
    row["选择项"] = _selection(row["仓库名称"], row["计费规则"])
    candidate = _normalize_database(pd.DataFrame([row]))
    combined = pd.concat([data, candidate], ignore_index=True)
    key = [
        "选择项",
        "生效日期",
        "品类",
        "天数下限",
        "天数上限",
    ]
    comparable = combined[key].fillna("<空>").astype(str)
    if comparable.duplicated().any():
        raise ValueError("相同选择项、版本、品类和天数档位的规则已存在")
    _write_database(combined, path)
    return candidate


def update_rate_rule(
    selection: str,
    effective_date: str | pd.Timestamp,
    updates: Mapping[str, Any],
    database_path: str | Path = DEFAULT_RATE_DATABASE,
    row_match: Mapping[str, Any] | None = None,
) -> int:
    """修改指定规则版本；只允许修改单价、生效日期、计费规则和备注。"""
    invalid = set(updates) - EDITABLE_FIELDS
    if invalid:
        raise ValueError(f"不允许修改字段：{', '.join(sorted(invalid))}")
    path = Path(database_path)
    data = load_rate_database(path)
    version = pd.Timestamp(effective_date).normalize()
    mask = (data["选择项"] == selection) & (data["生效日期"] == version)
    if row_match:
        for field, value in row_match.items():
            if field not in {"品类", "天数下限", "天数上限"}:
                raise ValueError(f"不支持的明细定位字段：{field}")
            if pd.isna(value):
                mask &= data[field].isna()
            else:
                mask &= data[field].astype(str) == str(value)
    if not mask.any():
        raise ValueError(f"未找到规则版本：{selection} / {version.date()}")
    for field, value in updates.items():
        data.loc[mask, field] = value
    if "计费规则" in updates:
        data.loc[mask, "选择项"] = [
            _selection(warehouse, updates["计费规则"])
            for warehouse in data.loc[mask, "仓库名称"]
        ]
    _write_database(data, path)
    return int(mask.sum())


def disable_rate_rule(
    selection: str,
    effective_date: str | pd.Timestamp | None = None,
    database_path: str | Path = DEFAULT_RATE_DATABASE,
) -> int:
    """停用规则或指定版本；保留全部历史数据。"""
    path = Path(database_path)
    data = load_rate_database(path)
    mask = data["选择项"] == selection
    if effective_date is not None:
        mask &= data["生效日期"] == pd.Timestamp(effective_date).normalize()
    if not mask.any():
        raise ValueError(f"未找到待停用规则：{selection}")
    data.loc[mask, "是否启用"] = "否"
    _write_database(data, path)
    return int(mask.sum())


def import_rate_config(
    source_path: str | Path,
    database_path: str | Path = DEFAULT_RATE_DATABASE,
    *,
    conflict: str = "error",
    new_effective_date: str | pd.Timestamp | None = None,
) -> dict[str, int]:
    """显式处理冲突后合并外部rate_config，不允许静默覆盖。"""
    if conflict not in {"error", "overwrite", "new_version"}:
        raise ValueError("conflict必须为error、overwrite或new_version")
    path = Path(database_path)
    source_frame = _read_rate_file(Path(source_path))
    imported = (
        _normalize_database(source_frame)
        if set(RATE_DATABASE_COLUMNS).issubset(source_frame.columns)
        else _legacy_to_database(source_frame)
    )
    if conflict == "new_version":
        if new_effective_date is None:
            raise ValueError("新增版本必须明确提供生效日期")
        imported["生效日期"] = pd.Timestamp(new_effective_date).normalize()
        imported["失效日期"] = pd.NaT
        imported["备注"] = (
            imported["备注"].astype(str) + "；作为新版本导入"
        )
    existing = (
        load_rate_database(path)
        if path.exists()
        else pd.DataFrame(columns=RATE_DATABASE_COLUMNS)
    )
    key = [
        "选择项",
        "生效日期",
        "品类",
        "天数下限",
        "天数上限",
    ]
    existing_keys = set(
        map(tuple, existing[key].fillna("<空>").astype(str).to_numpy())
    )
    imported_keys = list(
        map(tuple, imported[key].fillna("<空>").astype(str).to_numpy())
    )
    conflict_mask = pd.Series(
        [item in existing_keys for item in imported_keys],
        index=imported.index,
    )
    conflict_count = int(conflict_mask.sum())
    if conflict_count and conflict == "error":
        raise ValueError(
            f"发现{conflict_count}条已有规则；请选择覆盖或新增版本"
        )
    if conflict_count and conflict == "overwrite":
        conflict_keys = {
            imported_keys[index]
            for index, value in enumerate(conflict_mask)
            if value
        }
        keep = [
            tuple(row) not in conflict_keys
            for row in existing[key].fillna("<空>").astype(str).to_numpy()
        ]
        existing = existing.loc[keep]
    combined = pd.concat([existing, imported], ignore_index=True)
    _write_database(combined, path)
    return {
        "新增行数": int((~conflict_mask).sum()),
        "覆盖行数": conflict_count if conflict == "overwrite" else 0,
        "新增版本行数": len(imported) if conflict == "new_version" else 0,
    }


def diagnose_rate_rule_compatibility(
    selection: str,
    stat_date: str | pd.Timestamp,
    database_path: str | Path = DEFAULT_RATE_DATABASE,
) -> dict[str, object]:
    """识别冻结费用核心不能自动执行的费率类型，不参与费用计算。"""
    data = load_rate_database(database_path)
    date = pd.Timestamp(stat_date).normalize()
    rows = data[
        (data["选择项"] == str(selection).strip())
        & (data["是否启用"] == "是")
        & (data["生效日期"] <= date)
        & (data["失效日期"].isna() | (data["失效日期"] >= date))
    ].copy()
    if rows.empty:
        return {
            "诊断类型": "费率规则类型",
            "规则名称": str(selection).strip(),
            "当前规则类型": "未知",
            "可自动计算": False,
            "状态": "规则不可用",
            "为什么无法自动计算": "统计日期没有有效且启用的规则版本。",
            "建议处理方式": "请检查规则生效日期、失效日期和启用状态。",
        }
    version = rows["生效日期"].max()
    current = rows[rows["生效日期"] == version].copy()
    warehouse = str(current.iloc[0]["仓库名称"])
    has_tiers = current["天数上限"].notna().any() or (
        pd.to_numeric(current["天数下限"], errors="coerce").fillna(0) > 0
    ).any()
    if has_tiers:
        return {
            "诊断类型": "费率规则类型冲突",
            "仓库": warehouse,
            "规则名称": str(selection).strip(),
            "规则层级": "品类+存放天数分档",
            "当前规则类型": "存放天数分档",
            "规则版本": version.strftime("%Y-%m-%d"),
            "候选规则数量": int(len(current)),
            "可自动计算": False,
            "状态": "需要业务确认",
            "为什么无法自动计算": "当前冻结费用核心未启用存放天数分档计费模型。",
            "建议处理方式": (
                "A. 使用经业务确认的兼容统一单价规则继续计算；"
                "B. 返回维护费率规则。"
            ),
        }
    category_values = current["品类"].astype(str).str.strip()
    rule_type = "统一单价" if (category_values == "全部").all() else "品类单价"
    return {
        "诊断类型": "费率规则类型",
        "仓库": warehouse,
        "规则名称": str(selection).strip(),
        "规则层级": "仓库级" if rule_type == "统一单价" else "品类级",
        "当前规则类型": rule_type,
        "规则版本": version.strftime("%Y-%m-%d"),
        "候选规则数量": int(len(current)),
        "可自动计算": True,
        "状态": "兼容",
        "为什么无法自动计算": "",
        "建议处理方式": "可按现有冻结费用核心运行。",
    }


def get_effective_rate_rule(
    selection: str,
    stat_date: str | pd.Timestamp,
    database_path: str | Path = DEFAULT_RATE_DATABASE,
) -> tuple[pd.DataFrame, pd.Timestamp]:
    """选择统计日期有效的最近规则版本并转换为核心费率输入结构。"""
    data = load_rate_database(database_path)
    date = pd.Timestamp(stat_date).normalize()
    history = data[
        (data["选择项"] == selection)
        & (data["是否启用"] == "是")
        & (data["生效日期"] <= date)
    ].copy()
    if history.empty:
        raise ValueError(f"统计日期{date.date()}没有有效费率：{selection}")
    valid_at_stat = history[
        history["失效日期"].isna() | (history["失效日期"] >= date)
    ]
    if valid_at_stat.empty:
        raise ValueError(f"统计日期{date.date()}没有未失效费率：{selection}")
    version = valid_at_stat["生效日期"].max()
    compatibility = diagnose_rate_rule_compatibility(
        selection, date, database_path
    )
    if not compatibility["可自动计算"]:
        raise RateRuleCompatibilityError(compatibility)
    rates = pd.DataFrame(
        {
            "仓库名称": history["仓库名称"],
            "品类或物料": history["品类"],
            "单价": history["单价"],
            "计费方式说明": history["计价方式"],
            "适用层级": history["品类"].map(
                lambda value: "全部" if str(value).strip() == "全部" else "品类"
            ),
            "生效日期": history["生效日期"],
            "是否自动计算": "是",
            "特殊规则类型": "",
        }
    ).reset_index(drop=True)
    return rates, pd.Timestamp(version)
