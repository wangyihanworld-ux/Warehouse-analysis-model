from __future__ import annotations

from typing import Callable

from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

from warehouse_analysis.core.models import EngineResult


def append_attribution_sheet(
    wb,
    result: EngineResult,
    reported_test_percent: float | None,
    style_sheet: Callable,
    yellow: str,
) -> None:
    """写入仓储费差异归因工作表。

    输入：工作簿、引擎结果、测试上报比例及报告样式函数。
    输出：向工作簿追加与黄金版本一致的公式驱动归因Sheet。
    业务含义：分开显示上报静态差异与系统期间六因素变动。
    FIFO影响：仅消费汇总结果，不改变FIFO顺序、数量或仓储费。
    """
    attr = wb.create_sheet("仓储费差异归因分析")
    headers = [
        "仓库名称", "统计期间", "统计开始日", "统计截止日",
        "上期系统计算金额", "本期系统计算金额", "仓库实际上报仓储费金额",
        "上报静态差异", "系统费用期间变动",
        "本期已出库总量", "上期已出库总量", "本期已出库平均周转天数",
        "上期已出库平均周转天数", "本期已出库平均单价", "上期已出库平均单价",
        "已出库-数量贡献", "已出库-天数贡献", "已出库-单价贡献",
        "本期仍在库量", "上期仍在库量", "本期仍在库平均库龄",
        "上期仍在库平均库龄", "本期仍在库平均单价", "上期仍在库平均单价",
        "仍在库-数量贡献", "仍在库-天数贡献", "仍在库-单价贡献",
        "六项贡献合计", "期间归因交叉项", "交叉项占期间变动比例", "归因说明",
    ]
    for col, header in enumerate(headers, 1):
        attr.cell(1, col, header)
    totals = result.warehouse_total.sort_values(
        ["仓库名称", "统计截止日", "统计开始日"], kind="stable"
    ).reset_index(drop=True)
    prior_row_by_warehouse: dict[str, int] = {}
    for idx, total in totals.iterrows():
        row = idx + 2
        warehouse = total["仓库名称"]
        attr.cell(row, 1, warehouse)
        attr.cell(row, 2, total["统计期间"])
        attr.cell(row, 3, total["统计开始日"])
        attr.cell(row, 4, total["统计截止日"])
        attr.cell(row, 6, float(total["系统计算仓储费（元）"]))
        if reported_test_percent is not None:
            attr.cell(
                row,
                7,
                float(total["系统计算仓储费（元）"])
                * (1 + reported_test_percent),
            )
        attr.cell(row, 8, f'=IF(G{row}="","",G{row}-F{row})')
        attr.cell(row, 10, float(total["已出库总量（吨）"]))
        attr.cell(row, 12, float(total["已出库加权平均周转天数"]))
        attr.cell(row, 14, float(total["已出库平均单价（元/吨/天）"]))
        attr.cell(row, 19, float(total["当前总库存量（吨）"]))
        attr.cell(row, 21, float(total["仍在库加权平均库龄（天）"]))
        attr.cell(row, 23, float(total["仍在库平均单价（元/吨/天）"]))
        prior_row = prior_row_by_warehouse.get(warehouse)
        if prior_row is None:
            for col in (
                5, 9, 11, 13, 15, 16, 17, 18, 20, 22, 24, 25, 26,
                27, 28, 29, 30,
            ):
                attr.cell(row, col, None)
            attr.cell(
                row,
                31,
                f'=IF(G{row}="","",IF(H{row}=0,'
                '"系统计算金额与上报金额一致",'
                '"缺少上期数据，无法做期间归因，仅显示金额差异"))',
            )
        else:
            attr.cell(row, 5, f"=F{prior_row}")
            attr.cell(row, 9, f"=F{row}-E{row}")
            attr.cell(row, 11, f"=J{prior_row}")
            attr.cell(row, 13, f"=L{prior_row}")
            attr.cell(row, 15, f"=N{prior_row}")
            attr.cell(row, 16, f"=(J{row}-K{row})*M{row}*O{row}")
            attr.cell(row, 17, f"=J{row}*(L{row}-M{row})*O{row}")
            attr.cell(row, 18, f"=J{row}*L{row}*(N{row}-O{row})")
            attr.cell(row, 20, f"=S{prior_row}")
            attr.cell(row, 22, f"=U{prior_row}")
            attr.cell(row, 24, f"=W{prior_row}")
            attr.cell(row, 25, f"=(S{row}-T{row})*V{row}*X{row}")
            attr.cell(row, 26, f"=S{row}*(U{row}-V{row})*X{row}")
            attr.cell(row, 27, f"=S{row}*U{row}*(W{row}-X{row})")
            attr.cell(row, 28, f"=SUM(P{row}:R{row},Y{row}:AA{row})")
            attr.cell(row, 29, f"=I{row}-AB{row}")
            attr.cell(row, 30, f'=IF(I{row}=0,0,AC{row}/I{row})')
            for col, contribution in zip(
                range(32, 38), ("P", "Q", "R", "Y", "Z", "AA")
            ):
                attr.cell(row, col, f"=ABS({contribution}{row})")
            attr.cell(
                row, 38, f"=MATCH(MAX(AF{row}:AK{row}),AF{row}:AK{row},0)"
            )
            attr.cell(
                row,
                39,
                f"=MATCH(LARGE(AF{row}:AK{row},2),AF{row}:AK{row},0)",
            )
            attr.cell(
                row,
                31,
                f'=IF(I{row}=0,"本期与上期系统计算金额一致",'
                f'"系统费用期间变动主要由"&CHOOSE(AL{row},"已出库数量变化","已出库天数变化","已出库单价变化","仍在库数量变化","仍在库天数变化","仍在库单价变化")'
                f'&"引起（占期间变动"&TEXT(CHOOSE(AL{row},P{row},Q{row},R{row},Y{row},Z{row},AA{row})/I{row},"0.0%")'
                f'&"），其次是"&CHOOSE(AM{row},"已出库数量变化","已出库天数变化","已出库单价变化","仍在库数量变化","仍在库天数变化","仍在库单价变化")'
                f'&"（占期间变动"&TEXT(CHOOSE(AM{row},P{row},Q{row},R{row},Y{row},Z{row},AA{row})/I{row},"0.0%")&"）")',
            )
        prior_row_by_warehouse[warehouse] = row
    for col in range(32, 40):
        attr.column_dimensions[get_column_letter(col)].hidden = True
    style_sheet(attr, 1, attr.max_row, 31)
    attr.column_dimensions["G"].width = 26
    attr.column_dimensions["AE"].width = 70
    for row in range(2, attr.max_row + 1):
        attr.cell(row, 7).fill = PatternFill("solid", fgColor=yellow)
        for col in range(5, 31):
            if col != 31:
                attr.cell(row, col).number_format = "#,##0.000"
        attr.cell(row, 30).number_format = "0.0%"
