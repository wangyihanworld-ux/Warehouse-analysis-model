from __future__ import annotations

from pathlib import Path
from typing import Mapping, Sequence

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from warehouse_analysis.reporting.file_naming import ensure_unique_report_path


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _write_frame(sheet, frame: pd.DataFrame) -> None:
    headers = list(frame.columns)
    sheet.append(headers)
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    for row in frame.itertuples(index=False, name=None):
        sheet.append(list(row))
    sheet.freeze_panes = "A2"
    if headers:
        sheet.auto_filter.ref = sheet.dimensions
    for index, header in enumerate(headers, start=1):
        values = [str(header)] + [
            str(sheet.cell(row=row, column=index).value or "")
            for row in range(2, min(sheet.max_row, 200) + 1)
        ]
        sheet.column_dimensions[get_column_letter(index)].width = min(
            max(map(len, values)) + 3, 36
        )


def build_company_report(
    company_summary: pd.DataFrame,
    failures: Sequence[Mapping[str, object]],
    output_path: str | Path,
) -> Path:
    """输出企业仓储汇总报告，不改变任何单仓计算结果。"""
    destination = ensure_unique_report_path(Path(output_path))
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "企业仓储汇总"
    _write_frame(summary_sheet, company_summary)
    for row in range(2, summary_sheet.max_row + 1):
        for column in (3, 4):
            summary_sheet.cell(row, column).number_format = "#,##0.000"
        summary_sheet.cell(row, 5).number_format = "#,##0.00"
        summary_sheet.cell(row, 6).number_format = "#,##0.0000"
        summary_sheet.cell(row, 7).number_format = "#,##0.00"

    failure_frame = pd.DataFrame(
        [
            {
                "仓库名称": item.get("warehouse", ""),
                "统计期间": item.get("period", ""),
                "状态": item.get("status", "failed"),
                "错误原因": item.get("error", ""),
                "解决建议": item.get("suggestion", ""),
            }
            for item in failures
        ],
        columns=["仓库名称", "统计期间", "状态", "错误原因", "解决建议"],
    )
    failure_sheet = workbook.create_sheet("失败任务")
    _write_frame(failure_sheet, failure_frame)
    workbook.save(destination)
    return destination
