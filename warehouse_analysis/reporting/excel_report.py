from __future__ import annotations

from pathlib import Path
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from warehouse_analysis.analysis.attribution import append_attribution_sheet
from warehouse_analysis.core.models import EngineResult
from warehouse_analysis.reporting.file_naming import resolve_report_output_path
from warehouse_analysis.version import VERSION, version_label


BLUE = "4472C4"
LIGHT_BLUE = "D9EAF7"
YELLOW = "FFF2CC"
WHITE = "FFFFFF"
GRID = Side(style="thin", color="B7C9D6")
RED_FILL = PatternFill("solid", fgColor="F4CCCC")
GREEN_FILL = PatternFill("solid", fgColor="D9EAD3")

BUSINESS_REPORT_SHEETS = [
    "首页摘要",
    "仓库经营分析",
    "批次明细",
    "按物料汇总",
    "按品类汇总",
    "仓库总计",
    "月末库存快照",
    "费用敏感性分析",
    "敏感性影响排名",
    "费率匹配记录",
    "异常记录",
    "数量守恒校验",
]

HOME_COLUMNS = [
    "仓库名称",
    "统计期间",
    "期初库存",
    "期间入库量",
    "期间出库量",
    "期末库存",
    "平均库存",
    "平均存放天数",
    "平均费率",
    "仓储费",
    "异常数量",
    "人工核对数量",
    "数量守恒状态",
]

RATE_MATCH_COLUMNS = [
    "仓库",
    "物料/品类",
    "匹配层级",
    "匹配对象",
    "单价",
    "生效日期",
    "是否人工核对",
]


def _write_dataframe(
    ws, frame: pd.DataFrame, start_row: int = 1
) -> None:
    """写入DataFrame；只改变报告载体，不影响FIFO结果。"""
    for col, header in enumerate(frame.columns, 1):
        ws.cell(start_row, col, header)
    for row_idx, row in enumerate(
        frame.itertuples(index=False, name=None), start_row + 1
    ):
        for col_idx, value in enumerate(row, 1):
            ws.cell(row_idx, col_idx, value)


def _style_sheet(ws, header_row: int, max_row: int, max_col: int) -> None:
    """应用黄金版本工作表样式，不影响业务数据和FIFO结果。"""
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = (
        f"A{header_row}:{get_column_letter(max_col)}{max_row}"
        if max_row >= header_row
        else None
    )
    for cell in ws[header_row]:
        if cell.column <= max_col:
            cell.fill = PatternFill("solid", fgColor=BLUE)
            cell.font = Font(name="微软雅黑", bold=True, color=WHITE)
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            cell.border = Border(bottom=GRID)
    ws.row_dimensions[header_row].height = 38
    for row in range(header_row + 1, max_row + 1):
        for col in range(1, max_col + 1):
            cell = ws.cell(row, col)
            cell.font = Font(name="微软雅黑", size=10)
            cell.border = Border(bottom=GRID)
            if row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    for col in range(1, max_col + 1):
        values = [
            str(ws.cell(row, col).value or "")
            for row in range(
                header_row, min(max_row, header_row + 100) + 1
            )
        ]
        ws.column_dimensions[get_column_letter(col)].width = min(
            max(max(map(len, values), default=8) + 2, 12), 32
        )


def export_result(
    result: EngineResult,
    output_path: Path,
    reported_test_percent: float | None = None,
) -> Path:
    """导出与黄金版本结构一致的Excel。

    输入：EngineResult、输出路径和可选测试上报比例。
    输出：原有八张工作表及相同公式、样式和数字格式的xlsx。
    业务含义：将现有计算和归因结果交付给业务用户。
    FIFO影响：只读引擎结果，不改变FIFO计算结果。
    """
    wb = Workbook()
    wb.remove(wb.active)
    sheets = [
        ("批次明细", result.batch_details),
        ("按物料汇总", result.material_summary),
        ("按品类汇总", result.category_summary),
        ("仓库总计", result.warehouse_total),
        ("月末库存快照", result.monthly_snapshot),
        ("数量守恒校验", result.conservation),
        ("异常记录", result.anomalies),
    ]
    for name, frame in sheets:
        ws = wb.create_sheet(name)
        _write_dataframe(ws, frame)
        _style_sheet(ws, 1, ws.max_row, ws.max_column)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                header = ws.cell(1, cell.column).value
                if isinstance(header, str) and (
                    "日期" in header or header.endswith("日")
                ):
                    cell.number_format = "yyyy-mm-dd"
                elif (
                    isinstance(cell.value, (int, float))
                    and isinstance(header, str)
                    and (
                        "数量" in header
                        or "库存" in header
                        or "仓储费" in header
                        or "单价" in header
                        or "天数" in header
                        or "差额" in header
                    )
                ):
                    cell.number_format = "#,##0.000"

    append_attribution_sheet(
        wb,
        result,
        reported_test_percent,
        _style_sheet,
        YELLOW,
    )
    attr = wb["仓储费差异归因分析"]
    attr["G1"].comment = Comment(
        "黄色单元格为手工输入。修改后“上报静态差异”自动更新；"
        "期间因素归因以本期系统金额减上期系统金额为目标。",
        "Codex",
    )
    attr["I1"].comment = Comment(
        "六项连环替代贡献解释的是系统费用期间变动，与上报静态差异"
        "分开列示，避免混用两个比较基准。",
        "Codex",
    )
    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        wb.calculation.calcMode = "auto"
    except AttributeError:
        pass
    warehouse_name = "未命名仓库"
    start_date = pd.Timestamp.today().normalize()
    end_date = start_date
    if not result.warehouse_total.empty:
        total = result.warehouse_total.iloc[0]
        warehouse_name = str(total.get("仓库名称", warehouse_name))
        start_date = pd.Timestamp(total.get("统计开始日期", start_date))
        end_date = pd.Timestamp(total.get("统计截止日期", end_date))
    resolved_path = resolve_report_output_path(
        output_path,
        warehouse_name,
        start_date,
        end_date,
        VERSION,
    )
    resolved_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(resolved_path)
    return resolved_path


def _is_yes(value: object) -> bool:
    """判断业务布尔值是否表示“是”；仅用于报告筛选。"""
    return str(value).strip().lower() in {"是", "yes", "true", "1", "y"}


def _business_number_format(header: object) -> str | None:
    """按业务字段语义返回Excel数字格式，不改变底层数值。"""
    if not isinstance(header, str):
        return None
    if "日期" in header or header.endswith("日"):
        return "yyyy-mm-dd"
    if "比例" in header or "占比" in header:
        return "0.00%"
    if "费率" in header or "单价" in header:
        return "#,##0.0000"
    if "天数" in header or "库龄" in header:
        return "#,##0.00"
    if (
        "仓储费" in header
        or "费用" in header
        or "金额" in header
        or header.endswith("（元）")
    ):
        return "#,##0.00"
    if (
        "数量" in header
        or "库存" in header
        or "总量" in header
        or "（吨）" in header
    ):
        return "#,##0.000"
    return None


def _format_business_sheet(
    ws, header_row: int = 1, freeze_row: int | None = None
) -> None:
    """统一业务报告工作表样式、筛选、冻结窗格和数字格式。"""
    max_row = ws.max_row
    max_col = ws.max_column
    _style_sheet(ws, header_row, max_row, max_col)
    ws.freeze_panes = f"A{freeze_row or header_row + 1}"
    for col in range(1, max_col + 1):
        number_format = _business_number_format(ws.cell(header_row, col).value)
        if not number_format:
            continue
        for row in range(header_row + 1, max_row + 1):
            ws.cell(row, col).number_format = number_format


def _period_filter(
    frame: pd.DataFrame, summary: pd.Series
) -> pd.DataFrame:
    """按仓库和可用统计期间字段筛选报告结果行。"""
    if frame.empty:
        return frame
    mask = pd.Series(True, index=frame.index)
    warehouse_column = next(
        (
            column
            for column in ("仓库名称", "仓库")
            if column in frame.columns
        ),
        None,
    )
    if warehouse_column:
        mask &= frame[warehouse_column].astype(str) == str(
            summary["仓库名称"]
        )
    if "统计期间" in frame.columns:
        mask &= frame["统计期间"].astype(str) == str(summary["统计期间"])
    date_pairs = [
        ("统计开始日", "统计开始日期"),
        ("统计截止日", "统计截止日期"),
        ("统计开始日期", "统计开始日期"),
        ("统计截止日期", "统计截止日期"),
    ]
    for frame_column, summary_column in date_pairs:
        if frame_column in frame.columns:
            mask &= pd.to_datetime(frame[frame_column]) == pd.Timestamp(
                summary[summary_column]
            )
    return frame.loc[mask]


def _build_home_summary(
    result: EngineResult,
    warehouse_summary: pd.DataFrame,
    rate_match_result: pd.DataFrame,
) -> pd.DataFrame:
    """构建管理层首页摘要；只组合已有结果，不重新计算FIFO或费用。"""
    rows: list[dict[str, object]] = []
    for _, summary in warehouse_summary.iterrows():
        anomalies = _period_filter(result.anomalies, summary)
        rate_matches = _period_filter(rate_match_result, summary)
        manual_count = (
            int(rate_matches["是否人工核对"].map(_is_yes).sum())
            if "是否人工核对" in rate_matches.columns
            else 0
        )
        conservation = _period_filter(result.conservation, summary)
        unbalanced = (
            int((conservation["校验结果"] != "平衡").sum())
            if "校验结果" in conservation.columns
            else 0
        )
        rows.append(
            {
                "仓库名称": summary["仓库名称"],
                "统计期间": summary["统计期间"],
                "期初库存": summary["期初库存量（吨）"],
                "期间入库量": summary["期间入库量（吨）"],
                "期间出库量": summary["期间出库量（吨）"],
                "期末库存": summary["期末库存量（吨）"],
                "平均库存": summary["平均库存量（吨）"],
                "平均存放天数": summary["平均存放天数"],
                "平均费率": summary["平均仓储费率（元/吨/天）"],
                "仓储费": summary["仓储费总额（元）"],
                "异常数量": len(anomalies),
                "人工核对数量": manual_count,
                "数量守恒状态": (
                    "正常" if unbalanced == 0 else f"异常（{unbalanced}）"
                ),
            }
        )
    return pd.DataFrame(rows, columns=HOME_COLUMNS)


def _build_anomaly_report(
    result: EngineResult,
    rate_match_result: pd.DataFrame,
) -> pd.DataFrame:
    """合并FIFO异常和费率人工核对事项，形成统一异常工作表。"""
    frames: list[pd.DataFrame] = []
    if not result.anomalies.empty:
        fifo = result.anomalies.copy()
        fifo.insert(0, "异常类型", "FIFO异常")
        frames.append(fifo)
    if (
        not rate_match_result.empty
        and "是否人工核对" in rate_match_result.columns
    ):
        manual = rate_match_result[
            rate_match_result["是否人工核对"].map(_is_yes)
        ].copy()
        if not manual.empty:
            manual.insert(0, "异常类型", "人工核对事项")
            if "异常原因" not in manual.columns:
                source = next(
                    (
                        column
                        for column in (
                            "人工核对原因",
                            "特殊规则类型",
                            "计费方式说明",
                        )
                        if column in manual.columns
                    ),
                    None,
                )
                manual["异常原因"] = (
                    manual[source] if source else "费率规则要求人工核对"
                )
            frames.append(manual)
    if not frames:
        return pd.DataFrame(columns=["异常类型", "异常原因"])
    return pd.concat(frames, ignore_index=True, sort=False)


def build_business_report(
    result: EngineResult,
    warehouse_summary: pd.DataFrame,
    fee_sensitivity: pd.DataFrame,
    sensitivity_ranking: pd.DataFrame,
    attribution_result: pd.DataFrame | None,
    rate_match_result: pd.DataFrame | None,
    output_path: Path,
    report_context: dict[str, object] | None = None,
) -> Path:
    """生成面向业务人员的十二工作表仓储分析报告。

    输入：
    EngineResult、仓库经营分析、费用敏感性明细、敏感性排名、归因结果、
    费率匹配记录和输出路径。

    输出：
    固定Sheet顺序、统一格式且可由Excel重新打开的xlsx工作簿。

    业务含义：
    将计算、经营分析、敏感性和追溯结果组合成管理层与业务人员共用报告。
    attribution_result保留在服务接口中供后续报告版本使用；本阶段固定的
    十二张Sheet不新增归因页，旧export_result仍完整保留原归因功能。

    FIFO影响：
    本函数只消费传入DataFrame，不调用或修改FIFO、费率匹配及计算逻辑。
    """
    del attribution_result
    rate_matches = (
        rate_match_result.copy()
        if rate_match_result is not None
        else pd.DataFrame(columns=RATE_MATCH_COLUMNS)
    )
    home_summary = _build_home_summary(
        result, warehouse_summary, rate_matches
    )
    anomaly_report = _build_anomaly_report(result, rate_matches)

    wb = Workbook()
    wb.remove(wb.active)

    home = wb.create_sheet("首页摘要")
    last_column = get_column_letter(len(HOME_COLUMNS))
    home.merge_cells(f"A1:{last_column}1")
    home["A1"] = "仓储分析报告"
    home["A1"].font = Font(name="微软雅黑", size=18, bold=True, color=WHITE)
    home["A1"].fill = PatternFill("solid", fgColor=BLUE)
    home["A1"].alignment = Alignment(horizontal="center", vertical="center")
    home.row_dimensions[1].height = 30
    home.merge_cells(f"A2:{last_column}2")
    context = report_context or {}
    opening_description = ""
    if context.get("期初库存来源"):
        opening_description += (
            f" | 期初库存来源：{context['期初库存来源']}"
        )
    if context.get("日期口径"):
        opening_description += f" | 日期口径：{context['日期口径']}"
    home["A2"] = (
        "计算方式：FIFO | 库存单位：吨 | 费用单位：元 | "
        f"敏感性：单因素线性模拟{opening_description}"
    )
    home["A2"].font = Font(name="微软雅黑", italic=True, color="555555")
    home["A2"].alignment = Alignment(horizontal="left")
    home.merge_cells(f"A3:{last_column}3")
    rate_description = ""
    if not rate_matches.empty:
        selections = sorted(
            {
                f"{row.get('仓库', '')}+{row.get('计费规则', '')}"
                for _, row in rate_matches.iterrows()
                if str(row.get("计费规则", "")).strip()
            }
        )
        versions = sorted(
            {
                pd.Timestamp(value).strftime("%Y-%m-%d")
                for value in rate_matches.get("规则版本", pd.Series(dtype=object))
                if pd.notna(value)
            }
        )
        if selections:
            rate_description = (
                f" | 本次使用费率：{'、'.join(selections)}"
                f" | 费率版本日期：{'、'.join(versions)}"
            )
    home["A3"] = (
        f"模型名称及版本：{version_label()} | "
        f"报告生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        f"{rate_description}"
    )
    home["A3"].font = Font(name="微软雅黑", color="555555")
    home["A3"].alignment = Alignment(horizontal="left")
    _write_dataframe(home, home_summary, start_row=4)
    _format_business_sheet(home, header_row=4, freeze_row=5)
    home.auto_filter.ref = f"A4:{last_column}{home.max_row}"
    if home.max_row >= 5:
        home.conditional_formatting.add(
            f"L5:L{home.max_row}",
            CellIsRule(operator="greaterThan", formula=["0"], fill=RED_FILL),
        )
        home.conditional_formatting.add(
            f"M5:M{home.max_row}",
            FormulaRule(formula=['$M5<>"正常"'], fill=RED_FILL),
        )
        home.conditional_formatting.add(
            f"M5:M{home.max_row}",
            FormulaRule(formula=['$M5="正常"'], fill=GREEN_FILL),
        )

    report_frames = [
        ("仓库经营分析", warehouse_summary),
        ("批次明细", result.batch_details),
        ("按物料汇总", result.material_summary),
        ("按品类汇总", result.category_summary),
        ("仓库总计", result.warehouse_total),
        ("月末库存快照", result.monthly_snapshot),
        ("费用敏感性分析", fee_sensitivity),
        ("敏感性影响排名", sensitivity_ranking),
        ("费率匹配记录", rate_matches),
        ("异常记录", anomaly_report),
        ("数量守恒校验", result.conservation),
    ]
    for name, frame in report_frames:
        ws = wb.create_sheet(name)
        _write_dataframe(ws, frame)
        _format_business_sheet(ws)

    if wb.sheetnames != BUSINESS_REPORT_SHEETS:
        raise RuntimeError("业务报告Sheet顺序与固定规范不一致")
    warehouse_names = sorted(
        set(warehouse_summary["仓库名称"].astype(str))
    )
    warehouse_name = "、".join(warehouse_names) if warehouse_names else "未命名仓库"
    start_date = pd.to_datetime(
        warehouse_summary["统计开始日期"], errors="raise"
    ).min()
    end_date = pd.to_datetime(
        warehouse_summary["统计截止日期"], errors="raise"
    ).max()
    resolved_path = resolve_report_output_path(
        output_path,
        warehouse_name=warehouse_name,
        start_date=start_date,
        end_date=end_date,
        version=VERSION,
    )
    resolved_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(resolved_path)
    return resolved_path
